"""
backend/app/services/report_export_service.py

İstatistik & Raporlama V2 öncelik #3, Faz 3 madde F — "Dağıtım: Resend
e-posta + PDF/CSV/Excel export". Bu modül SADECE render mantığını içerir
(Supabase'e hiç dokunmaz, dependency-injection ile test edilebilir) —
zaten hesaplanmış bir rapor sözlüğünü (user_report_service.get_user_report,
organization_report_service.get_organization_report,
platform_snapshot_service.get_snapshots) ortak bir ReportDocument ara
temsiline çevirip CSV/XLSX/PDF byte'larına dönüştürür.

PDF'te Türkçe karakterler (İ, ş, ğ, ü, ö, ç) için reportlab'ın varsayılan
Helvetica fontu YETERSİZ (Type1, tam Unicode desteği yok) — bu yüzden
backend/app/assets/fonts/DejaVuSans(-Bold).ttf gömülü font olarak
kaydediliyor (Railway/nixpacks image'ında hangi sistem fontlarının olduğu
garanti değil, bu yüzden repo'ya gömüldü).
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font as XlsxFont
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)
from reportlab.lib.styles import ParagraphStyle

_FONTS_DIR = Path(__file__).resolve().parent.parent / "assets" / "fonts"
_FONT_REGISTERED = False


def _ensure_fonts_registered() -> None:
    global _FONT_REGISTERED
    if _FONT_REGISTERED:
        return
    pdfmetrics.registerFont(TTFont("DejaVuSans", str(_FONTS_DIR / "DejaVuSans.ttf")))
    pdfmetrics.registerFont(TTFont("DejaVuSans-Bold", str(_FONTS_DIR / "DejaVuSans-Bold.ttf")))
    _FONT_REGISTERED = True


@dataclass
class ReportSection:
    title: str
    kv_pairs: list[tuple[str, str]] = field(default_factory=list)
    table_headers: list[str] | None = None
    table_rows: list[list[str]] | None = None
    note: str | None = None


@dataclass
class ReportDocument:
    title: str
    subtitle: str
    generated_at: str
    sections: list[ReportSection] = field(default_factory=list)


def to_csv(doc: ReportDocument) -> bytes:
    """Basit, tek sayfalık düz CSV — her bölüm bir başlık satırı + veri
    satırlarıyla art arda yazılır (Excel/Sheets'te açılınca okunaklı)."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow([doc.title])
    writer.writerow([doc.subtitle])
    writer.writerow([f"Oluşturulma: {doc.generated_at}"])
    writer.writerow([])
    for section in doc.sections:
        writer.writerow([section.title])
        for label, value in section.kv_pairs:
            writer.writerow([label, value])
        if section.table_headers and section.table_rows is not None:
            writer.writerow(section.table_headers)
            for row in section.table_rows:
                writer.writerow(row)
        if section.note:
            writer.writerow([section.note])
        writer.writerow([])
    # BOM: Excel'in Türkçe karakterleri (İ, ş, ğ) doğru göstermesi için —
    # BOM'suz UTF-8 CSV'yi Excel çoğu zaman Latin-1 sanıp karakterleri bozuyor.
    return b"\xef\xbb\xbf" + buf.getvalue().encode("utf-8")


def to_xlsx(doc: ReportDocument) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Rapor"

    bold = XlsxFont(bold=True)
    title_font = XlsxFont(bold=True, size=14)

    row = 1
    ws.cell(row=row, column=1, value=doc.title).font = title_font
    row += 1
    ws.cell(row=row, column=1, value=doc.subtitle)
    row += 1
    ws.cell(row=row, column=1, value=f"Oluşturulma: {doc.generated_at}")
    row += 2

    for section in doc.sections:
        ws.cell(row=row, column=1, value=section.title).font = bold
        row += 1
        for label, value in section.kv_pairs:
            ws.cell(row=row, column=1, value=label)
            ws.cell(row=row, column=2, value=value)
            row += 1
        if section.table_headers and section.table_rows is not None:
            for col_idx, header in enumerate(section.table_headers, start=1):
                ws.cell(row=row, column=col_idx, value=header).font = bold
            row += 1
            for table_row in section.table_rows:
                for col_idx, cell_value in enumerate(table_row, start=1):
                    ws.cell(row=row, column=col_idx, value=cell_value)
                row += 1
        if section.note:
            ws.cell(row=row, column=1, value=section.note)
            row += 1
        row += 1

    # Sütun genişliklerini içeriğe göre kabaca ayarla (openpyxl otomatik yapmıyor).
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=8)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 10), 60)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(doc: ReportDocument) -> bytes:
    _ensure_fonts_registered()
    buf = io.BytesIO()
    pdf_doc = SimpleDocTemplate(
        buf, pagesize=A4,
        leftMargin=18 * mm, rightMargin=18 * mm, topMargin=16 * mm, bottomMargin=16 * mm,
    )

    title_style = ParagraphStyle("Title", fontName="DejaVuSans-Bold", fontSize=17, leading=21, spaceAfter=4)
    subtitle_style = ParagraphStyle("Subtitle", fontName="DejaVuSans", fontSize=10, leading=14, textColor=colors.HexColor("#64748B"))
    meta_style = ParagraphStyle("Meta", fontName="DejaVuSans", fontSize=8, leading=11, textColor=colors.HexColor("#94A3B8"), spaceAfter=10)
    section_style = ParagraphStyle("Section", fontName="DejaVuSans-Bold", fontSize=12, leading=15, spaceBefore=12, spaceAfter=6, textColor=colors.HexColor("#0F172A"))
    kv_style = ParagraphStyle("KV", fontName="DejaVuSans", fontSize=9.5, leading=13)
    note_style = ParagraphStyle("Note", fontName="DejaVuSans", fontSize=8.5, leading=12, textColor=colors.HexColor("#64748B"), spaceBefore=4)

    story = [
        Paragraph(doc.title, title_style),
        Paragraph(doc.subtitle, subtitle_style),
        Paragraph(f"Oluşturulma: {doc.generated_at}", meta_style),
    ]

    for section in doc.sections:
        story.append(Paragraph(section.title, section_style))
        for label, value in section.kv_pairs:
            story.append(Paragraph(f"<b>{label}:</b> {value}", kv_style))
        if section.table_headers and section.table_rows is not None:
            data = [section.table_headers] + section.table_rows
            table = Table(data, hAlign="LEFT", repeatRows=1)
            table.setStyle(TableStyle([
                ("FONTNAME", (0, 0), (-1, -1), "DejaVuSans"),
                ("FONTNAME", (0, 0), (-1, 0), "DejaVuSans-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8.5),
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEDFE")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#534AB7")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#E2E8F0")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ]))
            story.append(Spacer(1, 4))
            story.append(table)
        if section.note:
            story.append(Paragraph(section.note, note_style))

    pdf_doc.build(story)
    return buf.getvalue()


# ─────────────────────────────────────────────────────────────────────────
# "Builder" fonksiyonları — user_report_service/organization_report_service/
# platform_snapshot_service'in zaten hesapladığı dict'leri yukarıdaki ortak
# ReportDocument temsiline çevirir. Bu fonksiyonlar da Supabase'e dokunmaz —
# tamamen saf, girdi->çıktı dönüşümü (route katmanı gerçek dict'i çekip
# buraya veriyor).
#
# V1 kapsam kararı: export SADECE Türkçe (uygulamanın kendi rapor
# sayfalarındaki gibi 10 dilli i18n YOK) — export bir "indir/e-posta ile
# gönder" eklentisi, ekrandaki rapor sayfasının birebir çevirisi değil.
# İleride gerekirse reportLocale.ts/reportStrings.ts'deki gibi bir
# locale parametresi eklenebilir.
# ─────────────────────────────────────────────────────────────────────────

PERIOD_LABELS = {"week": "Haftalık", "month": "Aylık"}


def _fmt_num(value: int | float | None, suffix: str = "") -> str:
    return "—" if value is None else f"{value}{suffix}"


def _fmt_change(value: int | float | None) -> str:
    if value is None:
        return "yeni"
    sign = "+" if value >= 0 else ""
    return f"{sign}{value}%"


def _fmt_bool(value: bool | None) -> str:
    return "Evet" if value else "Hayır"


def build_user_report_document(
    report: dict[str, Any], *, username: str, generated_at: str
) -> ReportDocument:
    period_label = PERIOD_LABELS.get(report.get("period"), str(report.get("period", "")))
    study = report.get("study", {})
    streak = report.get("streak", {})
    vocab = report.get("vocabulary", {})
    games = report.get("games", {})
    exam = report.get("exam", {})
    quests = report.get("quests", {})
    badges = report.get("badges", {})
    league = report.get("league", {})
    subscription = report.get("subscription", {})
    platform = report.get("platform", {})

    topic_rows = [
        ["Zayıf: " + t["topic_tag"], str(t["attempts"]), f"%{t['accuracy']}"]
        for t in exam.get("weak_topics", [])
    ] + [
        ["Güçlü: " + t["topic_tag"], str(t["attempts"]), f"%{t['accuracy']}"]
        for t in exam.get("strong_topics", [])
    ]

    league_rows = [
        [
            h.get("week_start") or "—",
            h.get("tier_slug") or "—",
            h.get("outcome") or "—",
            str(h.get("final_rank")) if h.get("final_rank") is not None else "—",
            str(h.get("final_xp")) if h.get("final_xp") is not None else "—",
        ]
        for h in league.get("history", [])
    ]

    sections = [
        ReportSection(
            title="Çalışma Özeti",
            kv_pairs=[
                ("Çalışma Süresi (dakika)", f"{study.get('minutes_current', 0)} ({_fmt_change(study.get('minutes_change_pct'))})"),
                ("Oturum Sayısı", f"{study.get('sessions_current', 0)} ({_fmt_change(study.get('sessions_change_pct'))})"),
                ("Güncel Seri", f"{streak.get('current', 0)} gün"),
                ("En Uzun Seri", f"{streak.get('longest', 0)} gün"),
            ],
        ),
        ReportSection(
            title="Kelime Hazinesi",
            kv_pairs=[
                ("Toplam Kelime", str(vocab.get("total_words", 0))),
                ("Öğrenilen Kelime", f"{vocab.get('learned_words', 0)} (%{vocab.get('learned_pct', 0)})"),
                ("Bu Dönem Yeni Kelime", f"{vocab.get('new_words_current', 0)} ({_fmt_change(vocab.get('new_words_change_pct'))})"),
            ],
        ),
        ReportSection(
            title="Oyun Performansı",
            kv_pairs=[
                ("Oturum Sayısı", f"{games.get('sessions_current', 0)} ({_fmt_change(games.get('sessions_change_pct'))})"),
                ("Ortalama Skor", f"{games.get('avg_score_current', 0)} (önceki: {games.get('avg_score_previous', 0)})"),
            ],
        ),
        ReportSection(
            title="Sınav / Konu Doğruluğu",
            kv_pairs=[
                ("Doğruluk Oranı", _fmt_num(exam.get("accuracy_current"), "%")),
                ("Önceki Dönem", _fmt_num(exam.get("accuracy_previous"), "%")),
            ],
            table_headers=["Konu", "Deneme", "Doğruluk"] if topic_rows else None,
            table_rows=topic_rows or None,
        ),
        ReportSection(
            title="Görev Haritası",
            kv_pairs=[
                ("Toplam Tamamlanan", str(quests.get("completed_total", 0))),
                ("Bu Dönem Tamamlanan", str(quests.get("completed_current", 0))),
                ("İlerleme", f"%{quests.get('progress_pct', 0)}"),
            ],
        ),
        ReportSection(
            title="Rozetler",
            kv_pairs=[
                ("Toplam Kazanılan", str(badges.get("total_earned", 0))),
                ("Bu Dönem Kazanılan", str(badges.get("earned_current", 0))),
            ],
        ),
        ReportSection(
            title="Lig Durumu",
            kv_pairs=[("Güncel Kademe", league.get("current_tier") or "—")],
            table_headers=["Hafta", "Kademe", "Sonuç", "Sıralama", "XP"] if league_rows else None,
            table_rows=league_rows or None,
        ),
        ReportSection(
            title="Abonelik",
            kv_pairs=[
                ("Premium Durumu", _fmt_bool(subscription.get("is_premium"))),
                ("Premium Bitiş", subscription.get("premium_until") or "—"),
            ],
        ),
        ReportSection(
            title="Platform Karşılaştırması",
            kv_pairs=[
                ("Platform Ort. Süre (dk)", _fmt_num(platform.get("avg_minutes_current"))),
                ("Platform Ort. Yeni Kelime", _fmt_num(platform.get("avg_new_words_current"))),
                ("Platform Ort. Doğruluk", _fmt_num(platform.get("avg_accuracy_current"), "%")),
                ("XP Yüzdelik Dilimi", _fmt_num(platform.get("xp_percentile"), "%")),
            ],
            note=(
                "Bu karşılaştırma şu an aynı ülkedeki (Türkiye) kullanıcı tabanına dayanıyor."
                if platform.get("same_country_cohort")
                else None
            ),
        ),
    ]

    return ReportDocument(
        title=f"Lexis Kullanıcı Raporu — {username}",
        subtitle=f"{period_label} dönem — öğrenilen dil: {str(report.get('learning_lang', '—')).upper()}",
        generated_at=generated_at,
        sections=sections,
    )


def build_org_report_document(report: dict[str, Any], *, generated_at: str) -> ReportDocument:
    org = report.get("org", {})
    period_label = PERIOD_LABELS.get(report.get("period"), str(report.get("period", "")))
    study = report.get("study", {})
    accuracy = report.get("accuracy", {})
    vocab = report.get("vocabulary", {})
    top_learners = report.get("top_learners", [])
    weak_topics = report.get("weak_topics", [])

    sections = [
        ReportSection(
            title="Kurum Özeti",
            kv_pairs=[
                ("Kurum", org.get("name") or "—"),
                ("Plan", org.get("plan") or "—"),
                ("Üye Sayısı", str(report.get("member_count", 0))),
                ("Aktif Üye Sayısı", str(report.get("active_member_count", 0))),
            ],
        ),
        ReportSection(
            title="Çalışma Özeti",
            kv_pairs=[
                ("Toplam Süre (dakika)", f"{study.get('minutes_current', 0)} ({_fmt_change(study.get('minutes_change_pct'))})"),
                ("Oturum Sayısı", f"{study.get('sessions_current', 0)} ({_fmt_change(study.get('sessions_change_pct'))})"),
                ("Doğruluk Oranı", _fmt_num(accuracy.get("current"), "%")),
                ("Bu Dönem Yeni Kelime", f"{vocab.get('new_words_current', 0)} ({_fmt_change(vocab.get('new_words_change_pct'))})"),
                ("Bu Dönem Kazanılan Rozet", str(report.get("badges_earned_current", 0))),
            ],
        ),
        ReportSection(
            title="En Aktif Üyeler",
            table_headers=["Kullanıcı", "Kazanılan XP"] if top_learners else None,
            table_rows=[[t.get("username") or t.get("user_id", "—"), str(t.get("xp_gained", 0))] for t in top_learners] or None,
        ),
        ReportSection(
            title="Zayıf Konular",
            table_headers=["Konu", "Deneme", "Doğruluk"] if weak_topics else None,
            table_rows=[[t["topic_tag"], str(t["attempts"]), f"%{t['accuracy']}"] for t in weak_topics] or None,
        ),
    ]

    return ReportDocument(
        title=f"Lexis Kurum Raporu — {org.get('name') or '—'}",
        subtitle=f"{period_label} dönem",
        generated_at=generated_at,
        sections=sections,
    )


def build_platform_snapshots_document(
    snapshots: list[dict[str, Any]], *, generated_at: str
) -> ReportDocument:
    rows = [
        [
            str(s.get("snapshot_date", "—")),
            str(s.get("new_signups_count", 0)),
            str(s.get("active_users_count", 0)),
            str(s.get("total_study_minutes", 0)),
            str(s.get("total_new_words", 0)),
            _fmt_num(s.get("avg_topic_accuracy"), "%"),
            str(s.get("total_xp_awarded", 0)),
            str(s.get("premium_users_count", 0)),
            str(s.get("total_active_profiles", 0)),
        ]
        for s in snapshots
    ]
    return ReportDocument(
        title="Lexis Platform Günlük Özet Geçmişi",
        subtitle=f"Son {len(snapshots)} gün (bot hariç)",
        generated_at=generated_at,
        sections=[
            ReportSection(
                title="Günlük Özetler",
                table_headers=[
                    "Tarih", "Yeni Kayıt", "Aktif Kullanıcı", "Toplam Süre (dk)",
                    "Yeni Kelime", "Ort. Doğruluk", "Verilen XP", "Premium", "Toplam Aktif Profil",
                ],
                table_rows=rows,
                note="Ort. Doğruluk o gün topic_practice_attempts'e hiç kayıt düşmediyse '—' gösterir (0 ile karıştırılmamalı).",
            )
        ],
    )


_MEDIA_TYPES = {
    "csv": "text/csv; charset=utf-8",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "pdf": "application/pdf",
}

SUPPORTED_FORMATS = tuple(_MEDIA_TYPES.keys())


def render(doc: ReportDocument, fmt: str) -> tuple[bytes, str]:
    """(bytes, media_type) döner. Route katmanı bunu FastAPI Response'a
    sarmalıyor (bkz. stats.py/organizations.py/admin_platform.py::*_export_route).
    Desteklenmeyen fmt için ValueError — route bunu 400'e çeviriyor."""
    if fmt == "csv":
        return to_csv(doc), _MEDIA_TYPES["csv"]
    if fmt == "xlsx":
        return to_xlsx(doc), _MEDIA_TYPES["xlsx"]
    if fmt == "pdf":
        return to_pdf(doc), _MEDIA_TYPES["pdf"]
    raise ValueError(f"Desteklenmeyen export formatı: {fmt}")
